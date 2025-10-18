import os
import io
from PyPDF2 import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt


def create_watermark_pdf(text: str, tmp_path: str = "watermark_tmp.pdf") -> str:
    """
    Создаёт временный PDF-файл с водяным знаком (текстом).
    Возвращает путь к полученному файлу.
    """
    # Создаём PDF с текстом по центру страницы
    c = canvas.Canvas(tmp_path, pagesize=letter)
    width, height = letter
    c.saveState()
    c.setFont("Helvetica", 40)
    c.setFillColorRGB(0.6, 0.6, 0.6, alpha=0.3)
    c.translate(width / 2, height / 2)
    c.rotate(45)
    c.drawCentredString(0, 0, text)
    c.restoreState()
    c.save()
    return tmp_path


def add_pdf_watermark(input_pdf: str, output_pdf: str, watermark_text: str) -> None:
    """
    Накладывает водяной знак (текст) на каждую страницу PDF.
    """
    # Генерируем временный файл с водяным знаком
    tmp_watermark = create_watermark_pdf(watermark_text)

    # Читаем исходный и ватермарк-файлы
    original = PdfReader(input_pdf)
    watermark = PdfReader(tmp_watermark)
    watermark_page = watermark.pages[0]

    writer = PdfWriter()
    # Накладываем на каждую страницу
    for page in original.pages:
        page.merge_page(watermark_page)
        writer.add_page(page)

    # Сохраняем результат
    with open(output_pdf, "wb") as f:
        writer.write(f)

    # Удаляем временный файл
    try:
        os.remove(tmp_watermark)
    except OSError:
        pass


def add_docx_watermark(input_docx: str, output_docx: str, watermark_image: str, width: Inches = Inches(6)) -> None:
    """
    Добавляет изображение водяного знака в шапку каждого раздела DOCX-документа.
    watermark_image — путь к картинке-водяному знаку.
    width — ширина по размеру документа (по умолчанию 6 дюймов).
    """
    doc = Document(input_docx)
    for section in doc.sections:
        header = section.header
        # Вставляем картинку в первый параграф заголовка
        if header.is_linked_to_previous:
            header.is_linked_to_previous = False
        para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        run = para.add_run()
        run.add_picture(watermark_image, width=width)

    doc.save(output_docx)

def add_docx_text_watermark(input_docx: str, output_docx: str, text: str = "https://t.me/cryptomanevry") -> None:
    """
    Добавляет текстовый водяной знак в футер каждого раздела DOCX.
    """
    doc = Document(input_docx)
    for section in doc.sections:
        footer = section.footer
        # Если в футере есть параграф — берём его, иначе создаём новый
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = text
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        run = p.runs[0]
        run.font.size = Pt(8)
    doc.save(output_docx)

if __name__ == "__main__":
    # Пример использования
    # add_pdf_watermark("input.pdf", "output_watermarked.pdf", "Sample Watermark")
    # add_docx_watermark("input.docx", "output_watermarked.docx", "watermark.png")
    print("Watermark utilities loaded.")
